import os
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_doc():
    doc = Document()

    # Title
    title = doc.add_heading('Community Garden Management System', 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    subtitle = doc.add_paragraph('Project Architecture & Implementation Guide')
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # 1. Project Overview
    doc.add_heading('1. Project Overview', level=1)
    doc.add_paragraph(
        "The Community Garden Management System is a web application designed to fulfill the requirements "
        "of the CS251 Software Engineering – 1 module. It provides an intuitive portal for garden members, "
        "volunteers, and administrators to manage land plots, resources, and community activities."
    )
    doc.add_paragraph(
        "The core functionality includes:"
    )
    core_features = [
        "Land Module: Plot grid, lease management, soil health, and compost tracking.",
        "Resource Module: Tool library with check-in/out, consumable inventory, and seed bank.",
        "Volunteer Module: Task and shift scheduling, voting/polling, and incident reporting.",
        "Communication & Reporting: Custom email broadcasts, pest alerts, waitlist notifications, and data exporting (CSV/PDF)."
    ]
    for feature in core_features:
        doc.add_paragraph(feature, style='List Bullet')

    # 2. Technology Stack
    doc.add_heading('2. Technology Stack', level=1)
    doc.add_paragraph(
        "The project is built entirely from scratch using native web technologies to ensure a strong "
        "understanding of fundamental programming concepts:"
    )
    tech_stack = [
        "Backend: PHP 8+ (No external frameworks like Laravel are used).",
        "Frontend: HTML5, CSS3 (Vanilla CSS with CSS variables/tokens), Vanilla JavaScript (AJAX/Fetch API).",
        "Database: MySQL / MariaDB via XAMPP.",
        "Development Server: Local Apache Server via XAMPP."
    ]
    for tech in tech_stack:
        doc.add_paragraph(tech, style='List Bullet')

    # 3. Object-Oriented Programming (OOP) Implementation
    doc.add_heading('3. Object-Oriented Programming (OOP)', level=1)
    doc.add_paragraph(
        "The project strictly adheres to Object-Oriented Programming principles and the "
        "Model-View-Controller (MVC) architectural pattern. This separation of concerns improves "
        "maintainability, scalability, and code readability."
    )

    doc.add_heading('A. Singleton Pattern (Database)', level=2)
    doc.add_paragraph(
        "The database connection is managed using the Singleton Design Pattern. "
        "This guarantees that only one PDO connection to the MySQL database is open at any time, "
        "reducing memory overhead and preventing duplicate connections."
    )

    doc.add_heading('B. Model-View-Controller (MVC) Architecture', level=2)
    doc.add_paragraph(
        "The codebase is divided into three distinct layers. Base classes in the `core/` folder "
        "(`Model.php` and `Controller.php`) provide the foundation."
    )
    
    doc.add_heading('The Model Layer', level=3)
    doc.add_paragraph(
        "Models (e.g., `ToolModel`, `NotificationModel`) live in the `models/` folder. "
        "They inherit from the base `Model` class, which gives them direct access to the database. "
        "Models are exclusively responsible for data access: performing SQL queries, inserting records, "
        "and formatting database results."
    )

    doc.add_heading('The Controller Layer', level=3)
    doc.add_paragraph(
        "Controllers (e.g., `ToolController`, `NotificationController`) live in the `controllers/` folder. "
        "They act as the bridge between the View and the Model. They receive POST requests, validate form data, "
        "call the necessary methods on the Model, handle business logic (like checking permissions or "
        "generating alerts), and then render the correct View."
    )

    doc.add_heading('The View Layer', level=3)
    doc.add_paragraph(
        "Views (e.g., `tools_index.php`) live in the `views/` folder. They contain NO direct database queries. "
        "They receive data arrays from their Controller and use minimal PHP to iterate over that data, "
        "outputting clean, secure HTML markup."
    )

    # 4. How the Flow Works
    doc.add_heading('4. Request Lifecycle (How it works)', level=1)
    doc.add_paragraph(
        "When a user clicks a link or submits a form, the system follows a predictable path (Routing):"
    )

    steps = [
        "Router Execution: The user visits a URL like `/modules/resources/tools.php`.",
        "Controller Instantiation: The router file simply requires the `ToolController` and instantiates it.",
        "Logic & Validation: The Controller takes over, checking if the request was a form submission (POST) or a page load (GET).",
        "Data Fetching: The Controller asks the `ToolModel` for the required data (e.g., 'give me all available tools').",
        "View Rendering: The Controller passes the fetched data into the View, which renders the final webpage sent back to the browser."
    ]
    
    for i, step in enumerate(steps, 1):
        p = doc.add_paragraph(f"{i}. {step}")

    # 5. Advanced Features
    doc.add_heading('5. Advanced Technical Implementations', level=1)
    advanced = [
        "File Uploading: Uses `multipart/form-data` and strict MIME-type checking for the Media Manager and Tool Images.",
        "AJAX/Fetch APIs: The Plot Map uses JavaScript `fetch()` calls to ping database sensors in real-time without reloading the page.",
        "Role-Based Access Control (RBAC): Every controller method verifies if the user is an 'admin', 'warden', or regular member before executing.",
        "XSS & Security: Output is passed through a custom `e()` function which utilizes `htmlspecialchars()` to prevent Cross-Site Scripting."
    ]
    for adv in advanced:
        doc.add_paragraph(adv, style='List Bullet')

    # 6. Installation & Setup Instructions
    doc.add_heading('6. Installation & Setup Instructions', level=1)
    doc.add_paragraph(
        "To run this project locally, you will need XAMPP (or a similar AMP stack) installed on your machine. "
        "Follow these steps to get the application up and running:"
    )
    
    setup_steps = [
        "1. Install XAMPP: Download and install XAMPP for Windows/Mac/Linux.",
        "2. Start Services: Open the XAMPP Control Panel and start the 'Apache' and 'MySQL' modules.",
        "3. Clone the Repository: Copy the entire project folder (named 'garden') into the XAMPP htdocs directory. On Windows, this is typically located at 'C:\\xampp\\htdocs\\garden'.",
        "4. Setup the Database: Open your web browser and go to 'http://localhost/phpmyadmin'. Create a new database named 'garden_db'.",
        "5. Import Database Schema: In phpMyAdmin, select the 'garden_db' database, click the 'Import' tab, and upload the provided 'database.sql' (or similar SQL dump) file to structure the tables.",
        "6. Configure Connection: Open the project code in your IDE and navigate to 'config/constants.php' or 'config/db.php' to ensure the database credentials match your local setup (usually root as username and empty password).",
        "7. Access the Application: Open your web browser and navigate to 'http://localhost/garden'. You should see the login screen.",
        "8. Default Admin Login: Use the default administrator credentials provided by the team to log in and access the full system."
    ]
    
    for step in setup_steps:
        doc.add_paragraph(step)

    doc.add_paragraph("\n\nThis architecture ensures the application is highly modular, easily testable, and satisfies the CS251 software engineering rubrics for structured system design.")

    # Save document
    doc.save('Project_Documentation_v2.docx')
    print("Documentation generated successfully as Project_Documentation_v2.docx")

if __name__ == '__main__':
    create_doc()
